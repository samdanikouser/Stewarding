import csv
import io
import re
from datetime import datetime
from functools import partial

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import CondPageBreak, Flowable, Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .organization import stamp

FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe(value):
    """Stops spreadsheet apps from running text that starts like a formula."""
    if isinstance(value, str) and value[:1] in FORMULA_PREFIXES:
        return "'" + value
    return value


def _org_lines(org: dict | None) -> list[str]:
    """The organization's name and details, one line each, for the head of a report."""
    if not org or not org.get("name"):
        return []
    contact = " | ".join(part for part in (org.get("phone"), org.get("email"), org.get("website")) if part)
    legal = org.get("legal_name") or ""
    registration = f"Reg. {org['registration_number']}" if org.get("registration_number") else ""
    amounts = f"Amounts in {org['currency']}" if org.get("currency") else ""
    return [line for line in (org["name"], legal if legal != org["name"] else "", org.get("address") or "", contact, registration, amounts) if line]


def _generated(report: dict) -> str:
    tz = (report.get("organization") or {}).get("timezone") or "Asia/Kuwait"
    return stamp(datetime.fromisoformat(report["generated_at"].rstrip("Z")), tz)


def _header(report: dict) -> list[tuple[str, str]]:
    period = report["period"]
    return [("Report", f"{period['type'].title()} report"), ("Period", f"{period['label']} ({period['from']} to {period['to']})"),
            ("Locations", ", ".join(report["locations"])), ("Location time zones", ", ".join(report["timezones"])),
            ("Generated", _generated(report)), ("Prepared by", report.get("prepared_by") or "-"),
            ("Position", report.get("prepared_position") or "-")]


def to_csv(report: dict) -> bytes:
    buf = io.StringIO()
    out = csv.writer(buf)
    for line in _org_lines(report.get("organization")):
        out.writerow([safe(line)])
    for label, value in _header(report):
        out.writerow([label, safe(value)])
    for sec in report["sections"]:
        out.writerow([])
        out.writerow([safe(sec["title"])])
        out.writerow([safe(c["label"]) for c in sec["columns"]])
        for row in sec["rows"]:
            out.writerow([safe(row.get(c["key"])) for c in sec["columns"]])
        for total in sec["totals"]:
            out.writerow([safe(total["label"]), safe(total["value"])])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(report: dict) -> bytes:
    book = Workbook()
    info = book.active
    info.title = "Report"
    org_lines = _org_lines(report.get("organization"))
    for line in org_lines:
        info.append([safe(line)])
    if org_lines:
        info["A1"].font = Font(bold=True, size=13)
    for label, value in _header(report):
        info.append([label, safe(value)])
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70
    used = {"Report"}
    for index, sec in enumerate(report["sections"], 1):
        name = "".join(ch for ch in sec["title"] if ch not in "[]:*?/\\")[:28] or f"Section {index}"
        while name in used:
            name = f"{name[:25]}_{index}"
        used.add(name)
        sheet = book.create_sheet(name)
        sheet.append([safe(c["label"]) for c in sec["columns"]])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in sec["rows"]:
            sheet.append([safe(row.get(c["key"])) for c in sec["columns"]])
        for total in sec["totals"]:
            sheet.append([])
            sheet.append([safe(total["label"]), safe(total["value"])])
        for i in range(len(sec["columns"])):
            sheet.column_dimensions[chr(65 + i % 26)].width = 20
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def _esc(value) -> str:
    text = "" if value is None else str(value)
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


FILLS = {"ok": "DCFCE7", "pending": "FEF3C7", "fail": "FEE2E2"}
WIDE = {"corrective_action": 3, "task": 3, "comment": 2, "approved_by": 2, "check": 1.8, "verified_by": 1.5, "limits": 1.4,
        "measure": 1.5, "date": 1.3, "time": 0.8, "result": 0.9, "value": 1.1}


def _weights(group: dict) -> list[float]:
    """Column widths that follow the content: the longest unbreakable word, or most of a typical cell, within limits."""
    sample = group["rows"][:200]
    out = []
    for key, label in group["columns"]:
        texts = [str(_flat(row.get(key)) or "") for row in sample] or [""]
        longest = max(len(word) for text in [*texts, label] for word in (text.split() or [""]))
        typical = sorted(len(text) for text in texts)[int(len(texts) * 0.8)]
        whole = max(len(text) for text in texts)
        out.append(max(5.0, min(longest, 18) * 1.25, min(typical, 40) * 0.6, len(label) * 0.7, whole * 0.95 if whole <= 14 else 0))
    return out


SYMBOLS = {"ok": "\u2713", "pending": "\u2713?", "fail": "\u2717", "waived": "W", "na": "-"}
COLORS = {"ok": "#15803d", "pending": "#b45309", "fail": "#b91c1c"}


class Mark(Flowable):
    """A tick or a cross drawn as lines, so it prints the same in every viewer."""

    def __init__(self, state: str):
        super().__init__()
        self.state, self.width, self.height = state, 14, 9
        self.hAlign = "CENTER"

    def wrap(self, *_):
        return self.width, self.height

    def draw(self):
        canvas = self.canv
        canvas.setStrokeColor(colors.HexColor(COLORS[self.state]))
        canvas.setLineWidth(1.4)
        if self.state == "fail":
            canvas.line(4, 1, 10, 8)
            canvas.line(4, 8, 10, 1)
        else:
            path = canvas.beginPath()
            path.moveTo(3.5, 4.5)
            path.lineTo(6, 1.5)
            path.lineTo(11, 8)
            canvas.drawPath(path, stroke=1, fill=0)
            if self.state == "pending":
                canvas.setDash(1, 1)
                canvas.line(2, 0, 12, 0)


def _flat(value):
    """A chart cell as text for spreadsheets: a symbol plus the initials of who did it and who approved it."""
    if isinstance(value, dict):
        who = "/".join(part for part in (value.get("by"), value.get("ver")) if part)
        return f"{SYMBOLS[value['s']]} {who}".strip()
    return value


def _mark(key: str, value, style) -> object:
    if isinstance(value, dict):
        state = value["s"]
        center = ParagraphStyle("center", parent=style, alignment=1)
        if state in COLORS:
            who = "/".join(part for part in (value.get("by"), value.get("ver")) if part)
            return [Mark(state), Paragraph(_esc(who), ParagraphStyle("tiny", parent=center, fontSize=4.5, leading=5))]
        return Paragraph(f'<font color="#6b7280">{SYMBOLS[state]}</font>', center)
    if key == "value" and isinstance(value, str) and re.fullmatch(r"[0-9a-f]{32}\.(png|jpg|webp)", value):
        from ..api.uploads import upload_path

        path = upload_path(value)
        try:
            return Image(str(path), width=28 * mm, height=14 * mm, kind="proportional") if path.is_file() else Paragraph("Image not found", style)
        except OSError:
            return Paragraph("Image could not be read", style)
    if key in STATUS_KEYS and isinstance(value, str) and value:
        words = value.replace("_", " ").strip()
        shown = words.capitalize() if re.fullmatch(r"[a-z ]+", words) else words
        colour = "#15803d" if words.lower() in GOOD else "#b91c1c" if words.lower() in BAD else "#b45309" if words.lower() in WARN else None
        return Paragraph(f'<font color="{colour}"><b>{_esc(shown)}</b></font>' if colour else _esc(shown), style)
    return Paragraph(_esc(value), style)


def _record_meta(doc: dict) -> list[tuple[str, str]]:
    return [("Period", doc["period"]["label"]), ("Locations", ", ".join(doc["locations"]) or "-"),
            ("Prepared by", doc["prepared_by"]), ("Position", doc.get("prepared_position") or "-"), ("Generated", doc["generated_at"])]


def record_to_csv(doc: dict) -> bytes:
    buf = io.StringIO()
    out = csv.writer(buf)
    for line in _org_lines(doc.get("organization")):
        out.writerow([safe(line)])
    out.writerow([safe(doc["title"])])
    for label, value in _record_meta(doc):
        out.writerow([label, safe(value)])
    for group in doc["groups"]:
        out.writerow([])
        out.writerow([safe(group["title"])])
        out.writerow([safe(label) for _, label in group["columns"]])
        for row in group["rows"]:
            out.writerow([safe(_flat(row.get(key))) for key, _ in group["columns"]])
        for label, value in group["totals"]:
            out.writerow([label, value])
    return buf.getvalue().encode("utf-8-sig")


DECIMAL_TEXT = re.compile(r"^-?\d+\.(\d+)$")


def _number(value):
    """A figure written as text with decimals ("357.000", "-6.267") as a real number, shown with the same decimals, so Excel can add it up."""
    if isinstance(value, str) and (m := DECIMAL_TEXT.match(value)):
        return float(value), "#,##0." + "0" * len(m.group(1))
    return value, None


def record_to_xlsx(doc: dict) -> bytes:
    book = Workbook()
    book.remove(book.active)
    used: set[str] = set()
    for index, group in enumerate(doc["groups"] or [{"title": "No records", "meta": [], "columns": [], "rows": [], "totals": []}], 1):
        name = "".join(ch for ch in group["title"] if ch not in "[]:*?/\\")[:28] or f"Sheet {index}"
        while name in used:
            name = f"{name[:25]}_{index}"
        used.add(name)
        sheet = book.create_sheet(name)
        org_lines = _org_lines(doc.get("organization"))
        for line in org_lines:
            sheet.append([safe(line)])
        if org_lines:
            sheet["A1"].font = Font(bold=True, size=14)
        sheet.append([safe(doc["title"])])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True, size=13)
        sheet.append([safe(group["title"])])
        sheet.cell(row=sheet.max_row, column=1).font = Font(bold=True)
        for label, value in [*_record_meta(doc), *group["meta"]]:
            sheet.append([label, safe(value)])
        sheet.append([])
        header_row = sheet.max_row + 1
        sheet.append([safe(label) for _, label in group["columns"]])
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EEF0")
        for row in group["rows"]:
            cells = [_number(_flat(row.get(key))) for key, _ in group["columns"]]
            sheet.append([value if fmt else safe(value) for value, fmt in cells])
            for cell, (_, fmt) in zip(sheet[sheet.max_row], cells):
                if fmt:
                    cell.number_format = fmt
            if group.get("layout") == "chart":
                for cell in sheet[sheet.max_row]:
                    state = next((s for s, sym in SYMBOLS.items() if isinstance(cell.value, str) and cell.value.startswith(sym) and s != "na"), None)
                    if state in FILLS:
                        cell.fill = PatternFill("solid", fgColor=FILLS[state])
        sheet.append([])
        for label, value in group["totals"]:
            number, fmt = _number(value)
            sheet.append([label, number])
            if fmt:
                sheet.cell(row=sheet.max_row, column=2).number_format = fmt
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        for i, (key, _) in enumerate(group["columns"]):
            letter = sheet.cell(row=1, column=i + 1).column_letter
            sheet.column_dimensions[letter].width = 9 if group.get("layout") == "chart" and key != "task" else 16 * WIDE.get(key, 1)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


# ---- PDF layout, shared by the period reports and the records -------------------------------------------------------

PAGE = landscape(A4)
MARGIN = 28
BRAND = colors.HexColor("#1d4aa8")
BRAND_SOFT = colors.HexColor("#eaf1fd")
INK = colors.HexColor("#111827")
MUTED = colors.HexColor("#6b7280")
RULE = colors.HexColor("#d9dee5")
ZEBRA = colors.HexColor("#f6f8fb")
STATUS_KEYS = {"status", "state", "condition", "result"}
GOOD = {"completed", "verified", "present", "approved", "done", "pass", "ok", "match"}
BAD = {"failed", "overdue", "absent", "rejected", "fail", "voided", "not done", "missed"}
WARN = {"pending", "pending verification", "pending approval", "requested", "late", "reported", "in progress"}
NUMBER = re.compile(r"-?\d[\d,]*(\.\d+)?%?")


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()["BodyText"]

    def make(name: str, font: str = "Helvetica", colour=INK, **kw) -> ParagraphStyle:
        return ParagraphStyle(name, parent=base, fontName=font, textColor=colour, **kw)

    return {
        "title": make("title", "Helvetica-Bold", fontSize=22, leading=26),
        "subtitle": make("subtitle", colour=MUTED, fontSize=10.5, leading=14),
        "section": make("section", "Helvetica-Bold", BRAND, fontSize=12.5, leading=15),
        "label": make("label", colour=MUTED, fontSize=6.8, leading=8.5),
        "value": make("value", "Helvetica-Bold", fontSize=9, leading=11),
        "big": make("big", "Helvetica-Bold", BRAND, fontSize=13, leading=15),
        "cell": make("cell", fontSize=7.8, leading=9.6),
        "head": make("head", "Helvetica-Bold", colors.white, fontSize=7.6, leading=9.2),
        "small": make("small", colour=MUTED, fontSize=7.4, leading=9.5),
        "org": make("org", colour=MUTED, fontSize=8, leading=10),
        "orgname": make("orgname", "Helvetica-Bold", BRAND, fontSize=13, leading=16),
    }


def _masthead(org: dict | None, st: dict, width: float) -> list:
    """Logo on the left, organization name and details on the right; nothing when no organization is set."""
    lines = _org_lines(org)
    if not lines:
        return []
    text = [Paragraph(_esc(lines[0]), st["orgname"])] + [Paragraph(_esc(line), st["org"]) for line in lines[1:]]
    logo = None
    if org.get("logo_path"):
        w, h = ImageReader(org["logo_path"]).getSize()
        logo = Image(org["logo_path"], width=40 * w / h, height=40)
    left = (logo.drawWidth + 14) if logo else 0.1
    table = Table([[logo or "", text]], colWidths=[left, width - left])
    table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 1.2, BRAND),
                               ("BOTTOMPADDING", (0, 0), (-1, 0), 8), ("LEFTPADDING", (0, 0), (0, 0), 0)]))
    return [table, Spacer(1, 14)]


def _panel(st: dict, pairs: list[tuple[str, str]], width: float) -> Table:
    """Label over value, laid out in a few columns on a soft background."""
    per_row = len(pairs) if len(pairs) <= 5 else 3
    cells = [[Paragraph(_esc(label).upper(), st["label"]), Paragraph(_esc(value) or "-", st["value"])] for label, value in pairs]
    rows = [cells[i:i + per_row] + [""] * (per_row - len(cells[i:i + per_row])) for i in range(0, len(cells), per_row)]
    table = Table(rows, colWidths=[width / per_row] * per_row)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BRAND_SOFT), ("ROUNDEDCORNERS", [4, 4, 4, 4]), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 10), ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    return table


def _section(st: dict, title: str, meta: list, width: float) -> list:
    head = [Paragraph(_esc(title), st["section"])]
    if meta:
        head.append(Paragraph("   |   ".join(f"{_esc(label)}: <b>{_esc(value)}</b>" for label, value in meta), st["small"]))
    bar = Table([[head]], colWidths=[width])
    bar.setStyle(TableStyle([("LINEBEFORE", (0, 0), (0, 0), 3, BRAND), ("LEFTPADDING", (0, 0), (0, 0), 8), ("TOPPADDING", (0, 0), (0, 0), 1),
                             ("BOTTOMPADDING", (0, 0), (0, 0), 1)]))
    return [bar]


def _totals(st: dict, totals: list, width: float) -> Table:
    per_row = min(len(totals), 8)
    cells = [[Paragraph(_esc(label), st["label"]), Paragraph(_esc(value), st["big"])] for label, value in totals]
    rows = [cells[i:i + per_row] + [""] * (per_row - len(cells[i:i + per_row])) for i in range(0, len(cells), per_row)]
    table = Table(rows, colWidths=[width / per_row] * per_row)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), BRAND_SOFT), ("LINEABOVE", (0, 0), (-1, 0), 1.4, BRAND), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("LEFTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return table


def _numeric(key: str, rows: list[dict]) -> bool:
    values = [row.get(key) for row in rows[:200] if row.get(key) not in (None, "")]
    return bool(values) and all(not isinstance(v, dict) and NUMBER.fullmatch(str(v).strip()) for v in values)


def _data_table(st: dict, columns: list, rows: list[dict], widths: list[float], chart: bool, empty: str = "No records in this period") -> Table:
    numeric = [not chart and _numeric(key, rows) for key, _ in columns]
    body = [ParagraphStyle(f"n{i}", parent=st["cell"], alignment=2) if n else st["cell"] for i, n in enumerate(numeric)]
    heads = [Paragraph(_esc(label).replace("\n", "<br/>"), ParagraphStyle(f"h{i}", parent=st["head"], alignment=2 if numeric[i] else (1 if chart and key != "task" else 0)))
             for i, (key, label) in enumerate(columns)]
    data = [heads] + [[_mark(key, row.get(key), body[i]) for i, (key, _) in enumerate(columns)] for row in rows]
    if not rows:
        data.append([Paragraph(f"<i>{_esc(empty)}</i>", st["small"])] + [""] * (len(columns) - 1))
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [("BACKGROUND", (0, 0), (-1, 0), BRAND), ("VALIGN", (0, 0), (-1, -1), "MIDDLE" if chart else "TOP"),
             ("TOPPADDING", (0, 0), (-1, -1), 3.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5), ("LEFTPADDING", (0, 0), (-1, -1), 5),
             ("RIGHTPADDING", (0, 0), (-1, -1), 5), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ZEBRA]),
             ("LINEBELOW", (0, 1), (-1, -1), 0.3, RULE), ("LINEBELOW", (0, -1), (-1, -1), 0.8, BRAND)]
    if chart:
        style += [("INNERGRID", (0, 0), (-1, -1), 0.3, RULE), ("BOX", (0, 0), (-1, -1), 0.5, RULE)]
    if not rows:
        style.append(("SPAN", (0, 1), (-1, 1)))
    table.setStyle(TableStyle(style))
    return table


def _signature(st: dict, width: float, prepared_by: str | None) -> Table:
    gap = width * 0.04
    cols = [width * 0.3, gap, width * 0.3, gap, width * 0.2]
    labels = [Paragraph(f"Prepared by{(': ' + _esc(prepared_by)) if prepared_by else ''}", st["small"]), "", Paragraph("Reviewed and approved by", st["small"]), "", Paragraph("Date", st["small"])]
    table = Table([["", "", "", "", ""], labels], colWidths=cols, rowHeights=[30, None])
    table.setStyle(TableStyle([("LINEBELOW", (0, 0), (0, 0), 0.6, INK), ("LINEBELOW", (2, 0), (2, 0), 0.6, INK), ("LINEBELOW", (4, 0), (4, 0), 0.6, INK),
                               ("LEFTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 1), (-1, 1), 3)]))
    return table


class _Numbered(canvas.Canvas):
    """Draws 'Generated ... Page x of y' once the page count is known."""

    def __init__(self, *args, generated: str = "", **kwargs):
        super().__init__(*args, **kwargs)
        self._generated, self._states = generated, []

    def showPage(self):
        self._states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._states)
        for state in self._states:
            self.__dict__.update(state)
            self.setFont("Helvetica", 7)
            self.setFillColor(MUTED)
            self.drawRightString(PAGE[0] - MARGIN, 14, f"Generated {self._generated}     Page {self._pageNumber} of {total}")
            super().showPage()
        super().save()


def _furniture(org_name: str | None, title: str, note: str):
    """A colour strip on every page, a running header after the first, and a ruled footer."""

    def draw(canv, document):
        width, height = PAGE
        canv.saveState()
        canv.setFillColor(BRAND)
        canv.rect(0, height - 6, width, 6, stroke=0, fill=1)
        if document.page > 1:
            canv.setFont("Helvetica-Bold", 7.5)
            canv.setFillColor(INK)
            canv.drawString(MARGIN, height - 22, org_name or "")
            canv.setFont("Helvetica", 7.5)
            canv.setFillColor(MUTED)
            canv.drawRightString(width - MARGIN, height - 22, title)
            canv.setStrokeColor(RULE)
            canv.line(MARGIN, height - 27, width - MARGIN, height - 27)
        canv.setStrokeColor(RULE)
        canv.line(MARGIN, 25, width - MARGIN, 25)
        canv.setFont("Helvetica", 7)
        canv.setFillColor(MUTED)
        canv.drawString(MARGIN, 14, note[:120])
        canv.restoreState()

    return draw


def _render(*, title: str, subtitle: str, org: dict | None, meta: list, sections: list[dict], prepared_by: str | None,
            generated: str, note: str) -> bytes:
    st, width, org = _styles(), PAGE[0] - 2 * MARGIN, org or {}
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=PAGE, leftMargin=MARGIN, rightMargin=MARGIN, topMargin=40, bottomMargin=36, title=title,
                            author=prepared_by or "")
    story = _masthead(org, st, width) + [Paragraph(_esc(title), st["title"]), Paragraph(_esc(subtitle), st["subtitle"]), Spacer(1, 10), _panel(st, meta, width)]
    if not sections:
        story += [Spacer(1, 16), Paragraph("No records for this period.", st["small"])]
    for number, sec in enumerate(sections):
        if sec.get("page_break") and number:
            story.append(PageBreak())
        chart = sec.get("layout") == "chart"
        weights = [(8 if key == "task" else 1) if chart else w for (key, _), w in zip(sec["columns"], _weights(sec))]
        widths = [width * w / sum(weights) for w in weights]
        piece = [Spacer(1, 18), *_section(st, sec["title"], sec["meta"], width), Spacer(1, 7)]
        if sec.get("summary_first") and sec["totals"]:
            piece += [_totals(st, sec["totals"], width), Spacer(1, 9)]
        if sec.get("legend"):
            piece += [Paragraph(_esc(sec["legend"]), st["small"]), Spacer(1, 4)]
        piece.append(_data_table(st, sec["columns"], sec["rows"], widths, chart, sec.get("empty", "No records in this period")))
        if sec["totals"] and not sec.get("summary_first"):
            piece += [Spacer(1, 7), _totals(st, sec["totals"], width)]
        story += [KeepTogether(piece)] if len(sec["rows"]) <= 8 else [CondPageBreak(120), *piece]
        if sec.get("layout") == "sheet":
            story += [Spacer(1, 18), KeepTogether(_signature(st, width, prepared_by))]
    if not (sections and all(s.get("layout") == "sheet" for s in sections)):
        story += [Spacer(1, 26), KeepTogether(_signature(st, width, prepared_by))]
    decorate = _furniture(org.get("name"), title, note)
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate, canvasmaker=partial(_Numbered, generated=generated))
    return buf.getvalue()


def _footer_note(org: dict | None, title: str) -> str:
    org = org or {}
    return " - ".join(part for part in (org.get("name"), title, org.get("report_footer")) if part)


def to_pdf(report: dict) -> bytes:
    org = report.get("organization") or {}
    title = f"{report['period']['type'].title()} operations report"
    sections = [{"title": sec["title"], "meta": [], "columns": [(c["key"], c["label"]) for c in sec["columns"]], "rows": sec["rows"],
                 "totals": [(t["label"], t["value"]) for t in sec["totals"]]} for sec in report["sections"]]
    meta = [(label, value) for label, value in _header(report)[1:] if label != "Generated"]
    return _render(title=title, subtitle=report["period"]["label"], org=org, meta=meta, sections=sections, prepared_by=report.get("prepared_by"),
                   generated=_generated(report), note=_footer_note(org, title))


def record_to_pdf(doc: dict) -> bytes:
    org = doc.get("organization") or {}
    meta = [(label, value) for label, value in _record_meta(doc) if label not in ("Generated", "Period")]
    return _render(title=doc["title"], subtitle=doc["period"]["label"], org=org, meta=meta, sections=doc["groups"], prepared_by=doc.get("prepared_by"),
                   generated=str(doc["generated_at"]), note=_footer_note(org, f"{doc['title']} - {doc['period']['label']}"))
