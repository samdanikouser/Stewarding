import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import AuditLog

from .helpers import loc_id
from .test_approvals import SUP_RB, complete, verify
from .test_checklists import MGR, STEWARD, SUP, find, generate
from .test_monitoring import mk_stage, record
from .test_traceability import food, receive
from .test_two_step_logs import make

URL = "/api/reports/records"


def monitoring(client, h, **params):
    query = "&".join(f"{k}={v}" for k, v in {"date_from": "2026-03-02", "date_to": "2026-03-02", **params}.items())
    return client.get(f"{URL}/monitoring?{query}", headers=h)


def seed_readings(client, login):
    steward, sup = login(STEWARD), login(SUP)
    eq, wash = mk_stage(client, steward)
    _, rinse = mk_stage(client, steward, "final_rinse")
    good = record(client, steward, eq, wash, 60).json()
    record(client, steward, eq, rinse, 70, corrective_action="Descaled the heater").json()
    client.post("/api/monitoring/readings/review", json={"reading_ids": [good["id"]]}, headers=sup)


def test_monitoring_record_lists_who_what_when_and_the_corrective_action(client, login, db, clock):
    seed_readings(client, login)
    doc = monitoring(client, login(MGR), location_id=loc_id(db, "MK")).json()
    assert doc["title"] == "Monitoring Record" and doc["locations"] == ["Main Kitchen"] and doc["prepared_by"] == "Maya Manager" and doc["prepared_position"] == "Back of House Manager"
    group = next(g for g in doc["groups"] if "DW-MK-1" in g["title"])
    assert group["title"] == "Main Kitchen - Main Kitchen Dishwasher 1 (DW-MK-1)"
    assert dict(group["meta"])["Type"] == "Dishwasher"
    wash, rinse = group["rows"]
    assert (wash["check"], wash["value"], wash["limits"], wash["result"]) == ("Wash", "60 °C", "55 to 65 °C", "Pass")
    assert wash["date"] == "02 Mar 2026" and wash["time"] == "06:05" and wash["recorded_by"] == "Sid Steward"
    assert wash["approved_by"].startswith("Sam Supervisor 02 Mar 2026 06:05")
    assert rinse["result"] == "FAIL" and rinse["approved_by"] == "Awaiting approval"
    assert "Descaled the heater" in rinse["corrective_action"] and "Sid Steward" in rinse["corrective_action"]
    totals = dict(group["totals"])
    assert totals == {"Checks recorded": 2, "Passed": 1, "Failed": 1, "Awaiting approval": 1}


def test_records_export_as_pdf_excel_and_csv_and_are_audited(client, login, db, clock):
    seed_readings(client, login)
    h = login(MGR)
    pdf = client.get(f"{URL}/monitoring/export?fmt=pdf&date_from=2026-03-02&date_to=2026-03-02", headers=h)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF") and "attachment" in pdf.headers["content-disposition"]
    xlsx = client.get(f"{URL}/monitoring/export?fmt=xlsx&date_from=2026-03-02&date_to=2026-03-02", headers=h)
    cells = [c.value for sheet in load_workbook(io.BytesIO(xlsx.content)).worksheets for row in sheet.iter_rows() for c in row if c.value]
    assert "Monitoring Record" in cells and "Descaled the heater (Sid Steward, retest due)" in cells
    csv = client.get(f"{URL}/monitoring/export?fmt=csv&date_from=2026-03-02&date_to=2026-03-02", headers=h).content.decode("utf-8-sig")
    assert "Monitoring Record" in csv and "Recorded by" in csv and "Sid Steward" in csv
    assert len(db.scalars(select(AuditLog).where(AuditLog.action == "report.export", AuditLog.entity == "record")).all()) == 3


def test_empty_period_still_produces_a_document(client, login, db, clock):
    h = login(MGR)
    doc = monitoring(client, h, date_from="2026-01-01", date_to="2026-01-01").json()
    assert doc["groups"] and all(g["rows"] == [] for g in doc["groups"])
    pdf = client.get(f"{URL}/monitoring/export?fmt=pdf&date_from=2026-01-01&date_to=2026-01-01", headers=h)
    assert pdf.content.startswith(b"%PDF")


def test_record_scope_validation_and_permissions(client, login, db, clock):
    seed_readings(client, login)
    assert monitoring(client, login(STEWARD)).status_code == 403
    assert client.get(f"{URL}/monitoring/export", headers=login(STEWARD)).status_code == 403
    assert monitoring(client, login(SUP_RB), location_id=loc_id(db, "MK")).status_code == 403
    assert all(g["rows"] == [] for g in monitoring(client, login(SUP_RB)).json()["groups"])  # the rooftop bar has no readings
    mine = monitoring(client, login(SUP)).json()
    assert {g["meta"][0][1] for g in mine["groups"] if g["rows"]} == {"Main Kitchen"}
    assert monitoring(client, login(MGR), date_from="2026-03-05", date_to="2026-03-01").status_code == 400
    assert monitoring(client, login(MGR), date_from="2024-01-01", date_to="2026-03-01").status_code == 400
    assert client.get(f"{URL}/nonsense", headers=login(MGR)).status_code == 404
    assert client.get(f"{URL}/monitoring", headers=None).status_code == 401


def test_a_month_record_filters_by_monitoring_point(client, login, db, clock):
    seed_readings(client, login)
    h = login(MGR)
    eq, _ = mk_stage(client, h)
    other, _ = mk_stage(client, h, "wash", "GW-MK-1")
    doc = monitoring(client, h, date_from="2026-03-01", date_to="2026-03-31", equipment_id=eq["id"]).json()
    assert len(doc["groups"]) == 1 and doc["period"]["label"] == "01 Mar 2026 to 31 Mar 2026"
    quiet = monitoring(client, h, equipment_id=other["id"]).json()["groups"]
    assert len(quiet) == 1 and quiet[0]["rows"] == []


def test_checklist_record_shows_completion_and_approval(client, login, db, clock):
    generate(client, login, db)
    steward = login(STEWARD)
    done = find(client, steward, "Clean drains")
    complete(client, steward, done)
    verify(client, login(SUP), done)
    complete(client, steward, find(client, steward, "Separate and remove waste"))
    doc = client.get(f"{URL}/checklists?date_from=2026-03-02&date_to=2026-03-02&location_id={loc_id(db, 'MK')}", headers=login(MGR)).json()
    group = next(g for g in doc["groups"] if "daily" in g["title"])
    row = next(r for r in group["rows"] if r["task"] == "Clean drains and remove debris")
    assert (row["status"], row["completed_by"], row["verified_by"]) == ("verified", "Sid Steward", "Sam Supervisor")
    waiting = next(r for r in group["rows"] if r["task"] == "Separate and remove waste")
    assert waiting["status"] == "pending verification" and waiting["verified_by"] == ""
    totals = dict(group["totals"])
    assert totals["Approved"] == 1 and totals["Awaiting approval"] == 1
    pdf = client.get(f"{URL}/checklists/export?fmt=pdf&date_from=2026-03-02&date_to=2026-03-02", headers=login(MGR))
    assert pdf.content.startswith(b"%PDF")


def get(client, h, kind, **params):
    base = {"date_from": "2026-03-02", "date_to": "2026-03-02", **params}
    return client.get(f"{URL}/{kind}?" + "&".join(f"{k}={v}" for k, v in base.items() if v is not None), headers=h)


def test_checklist_chart_has_tasks_down_and_dates_across_with_ticks_and_initials(client, login, db, clock):
    generate(client, login, db)
    steward = login(STEWARD)
    done = find(client, steward, "Clean drains")
    complete(client, steward, done)
    verify(client, login(SUP), done)
    complete(client, steward, find(client, steward, "Separate and remove waste"))
    clock.set(2026, 3, 2, 10, 0)  # after the 13:30 deadline the untouched tasks are overdue
    doc = get(client, login(MGR), "checklist_chart", date_to="2026-03-04", location_id=loc_id(db, "MK")).json()
    chart = next(g for g in doc["groups"] if g.get("layout") == "chart" and "daily" in g["title"])
    assert [label for _, label in chart["columns"]][:4] == ["Task", "Mon\n2", "Tue\n3", "Wed\n4"]
    rows = {r["task"]: r for r in chart["rows"]}
    assert rows["Clean drains and remove debris"]["d0"] == {"s": "ok", "by": "SS", "ver": "SS2"} or rows["Clean drains and remove debris"]["d0"]["s"] == "ok"
    assert rows["Clean drains and remove debris"]["d1"] == {"s": "na"}
    assert rows["Separate and remove waste"]["d0"]["s"] == "pending"
    assert rows["Clean worktables and food preparation surfaces"]["d0"]["s"] == "fail"
    assert "SS = Sid Steward" in dict(chart["meta"])["Initials"]
    exceptions = next(g for g in doc["groups"] if "Exceptions" in g["title"])
    assert {"Not done"} == {r["status"] for r in exceptions["rows"]}
    for fmt in ("pdf", "xlsx", "csv"):
        out = client.get(f"{URL}/checklist_chart/export?fmt={fmt}&date_from=2026-03-02&date_to=2026-03-04", headers=login(MGR))
        assert out.status_code == 200 and len(out.content) > 500
    csv = client.get(f"{URL}/checklist_chart/export?fmt=csv&date_from=2026-03-02&date_to=2026-03-04", headers=login(MGR)).content.decode("utf-8-sig")
    assert "✓" in csv and "✗" in csv


def test_checklist_chart_is_limited_to_a_month_and_can_filter_by_frequency(client, login, db, clock):
    generate(client, login, db)
    h = login(MGR)
    assert get(client, h, "checklist_chart", date_from="2026-03-01", date_to="2026-04-15").status_code == 400
    weekly = get(client, h, "checklist_chart", frequency="weekly", location_id=loc_id(db, "MK")).json()
    assert all("weekly" in g["title"] for g in weekly["groups"] if g.get("layout") == "chart")
    assert get(client, h, "checklist_chart", frequency="monthly").status_code == 400


def test_attendance_record_shows_presence_absence_and_leave(client, login, db, clock):
    client.post("/api/attendance/check-in", json={"location_id": loc_id(db, "MK")}, headers=login(STEWARD))
    mgr = login(MGR)
    emp = next(e["id"] for e in client.get("/api/employees?size=50", headers=mgr).json()["items"] if e["employee_code"] == "E005")
    client.post("/api/absences", json={"employee_id": emp, "location_id": loc_id(db, "BQ"), "day": "2026-03-02", "absence_type": "sick_leave"}, headers=mgr)
    clock.advance(hours=10)
    doc = get(client, mgr, "attendance").json()
    rows = [r for g in doc["groups"] for r in g["rows"]]
    by = {r["employee"]: r for r in rows}
    assert by["Sid Steward"]["status"] in ("Present", "On duty", "Missed checkout") and by["Sid Steward"]["check_in"] == "06:05"
    assert by["Bina Steward"]["status"] == "Sick leave"
    mk = next(g for g in doc["groups"] if "Main Kitchen" in g["title"])
    assert dict(mk["totals"])["Scheduled"] >= 1
    only = get(client, mgr, "attendance", employee_id=emp).json()
    assert {r["employee"] for g in only["groups"] for r in g["rows"]} == {"Bina Steward"}
    assert client.get(f"{URL}/attendance/export?fmt=pdf&date_from=2026-03-02&date_to=2026-03-02", headers=mgr).content.startswith(b"%PDF")


def test_receipts_record_totals_discounts_and_supplier_filter(client, login, db, clock):
    from .test_receiving import receive, suppliers
    from .helpers import item_id
    h = login(SUP)
    lines = [{"item_id": item_id(db, "PL-001"), "qty": 10, "unit_cost": 5, "discount_type": "percent", "discount": 10}]
    assert receive(client, h, db, lines, invoice="R-1").status_code == 201
    other = receive(client, h, db, [{"item_id": item_id(db, "BW-001"), "qty": 4, "unit_cost": 2.5}], invoice="R-2", supplier="Hotel Ware Trading")
    assert other.status_code == 201
    mgr = login(MGR)
    doc = get(client, mgr, "receipts", location_id=loc_id(db, "MK")).json()
    group = doc["groups"][0]
    first = next(r for r in group["rows"] if r["invoice"] == "R-1")
    assert (first["supplier"], first["qty"], first["unit_cost"], first["discount"], first["line_total"], first["received_by"]) == (
        "Gulf Catering Supplies", 10, "5.000", "5.000", "45.000", "Sam Supervisor")
    totals = dict(group["totals"])
    assert totals["Documents"] == 2 and totals["Sub total"] == "60.000" and totals["Total discount"] == "5.000" and totals["Invoice total (posted receipts)"] == "55.000"
    gulf = suppliers(client)["Gulf Catering Supplies"]
    filtered = get(client, mgr, "receipts", supplier_id=gulf, location_id=loc_id(db, "MK")).json()
    assert {r["invoice"] for g in filtered["groups"] for r in g["rows"]} == {"R-1"}
    rec = next(r["id"] for r in client.get("/api/receipts", headers=mgr).json()["items"] if r["invoice_number"] == "R-1")
    client.post(f"/api/receipts/{rec}/void", json={"reason": "Entered by mistake"}, headers=mgr)
    after = dict(get(client, mgr, "receipts", location_id=loc_id(db, "MK")).json()["groups"][0]["totals"])
    assert after["Invoice total (posted receipts)"] == "10.000"


def test_movements_record_filters_by_type_and_totals_in_and_out(client, login, db, clock):
    from .helpers import item_id
    sup = login(SUP)
    base = {"item_id": item_id(db, "PL-001"), "location_id": loc_id(db, "MK")}
    client.post("/api/inventory/transactions", json={**base, "txn_type": "receipt", "qty": 30}, headers=sup)
    client.post("/api/inventory/transactions", json={**base, "txn_type": "issue", "qty": 8, "reason": "Banquet"}, headers=sup)
    mgr = login(MGR)
    issues = get(client, mgr, "movements", movement="issues", location_id=loc_id(db, "MK")).json()
    group = issues["groups"][0]
    assert [r["type"] for r in group["rows"]] == ["Issue"] and group["rows"][0]["qty"] == -8 and group["rows"][0]["reason"] == "Banquet"
    assert dict(group["totals"]) == {"Movements": 1, "Stock in": 0, "Stock out": 8, "Net": -8}
    everything = get(client, mgr, "movements", location_id=loc_id(db, "MK")).json()["groups"][0]
    assert any(r["type"] == "Opening" for r in everything["rows"])
    assert dict(everything["totals"])["Net"] == sum(r["qty"] for r in everything["rows"])
    assert get(client, mgr, "movements", movement="nonsense").status_code == 400
    one = get(client, mgr, "movements", item_id=item_id(db, "BW-001"), location_id=loc_id(db, "MK")).json()
    assert all("BW-001" in r["item"] for g in one["groups"] for r in g["rows"])


def test_attendance_sheet_has_one_page_per_employee_with_day_totals(client, login, db, clock):
    from .test_absences import emp_id, record as absence
    mgr = login(MGR)
    manual_body = {"employee_id": emp_id(client, mgr, "E004"), "location_id": loc_id(db, "MK"), "day": "2026-02-27",
                   "check_in_time": "06:00:00", "check_out_time": "14:00:00", "reason": "Forgot to punch, confirmed"}
    assert client.post("/api/attendance/manual", json=manual_body, headers=mgr).status_code == 201
    assert absence(client, mgr, db, day="2026-02-28", kind="absent").status_code == 201
    assert absence(client, mgr, db, day="2026-03-01", kind="day_off").status_code == 201
    client.post("/api/attendance/check-in", json={"location_id": loc_id(db, "MK")}, headers=login(STEWARD))
    doc = client.get(f"{URL}/attendance_sheet?date_from=2026-02-27&date_to=2026-03-02&employee_id={emp_id(client, mgr, 'E004')}", headers=mgr).json()
    assert doc["title"] == "Employee Attendance Sheet" and len(doc["groups"]) == 1
    sheet = doc["groups"][0]
    assert sheet["title"] == "Sid Steward (E004)" and sheet["page_break"] and dict(sheet["meta"])["Job title"] == "Steward"
    assert dict(sheet["totals"]) == {"Days present": 2, "Late arrivals": 0, "Week off": 1, "Absent": 1, "On leave": 0,
                                     "Hours worked": "8:00", "Overtime": "0:00"}
    assert [(r["date"], r["weekday"], r["status"]) for r in sheet["rows"]][:3] == [
        ("27 Feb 2026", "Fri", "Present"), ("28 Feb 2026", "Sat", "Absent"), ("01 Mar 2026", "Sun", "Week off")]
    everyone = client.get(f"{URL}/attendance_sheet?date_from=2026-02-27&date_to=2026-03-02", headers=mgr).json()
    names = [g["title"] for g in everyone["groups"]]
    assert names == sorted(names) and len(names) >= 3
    pdf = client.get(f"{URL}/attendance_sheet/export?fmt=pdf&date_from=2026-02-27&date_to=2026-03-02", headers=mgr)
    import re
    assert pdf.content.startswith(b"%PDF") and len(re.findall(rb"/Type\s*/Page[^s]", pdf.content)) >= len(names)
    xlsx = client.get(f"{URL}/attendance_sheet/export?fmt=xlsx&date_from=2026-02-27&date_to=2026-03-02", headers=mgr)
    assert len(load_workbook(io.BytesIO(xlsx.content)).worksheets) == len(names)


def test_the_report_header_gives_name_and_position_separately(client, login, db, clock):
    sup = login(SUP)
    q = "date_from=2026-03-02&date_to=2026-03-02"
    doc = client.get(f"{URL}/attendance?{q}", headers=sup).json()
    assert (doc["prepared_by"], doc["prepared_position"]) == ("Sam Supervisor", "Stewarding Supervisor")
    csv = client.get(f"{URL}/attendance/export?fmt=csv&{q}", headers=login(MGR)).content.decode("utf-8-sig").splitlines()
    assert "Prepared by,Maya Manager" in csv and "Position,Back of House Manager" in csv
    sheet = load_workbook(io.BytesIO(client.get(f"{URL}/attendance/export?fmt=xlsx&{q}", headers=login(MGR)).content)).worksheets[0]
    cells = {(c.value, sheet.cell(row=c.row, column=2).value) for row in sheet.iter_rows(max_col=1) for c in row if c.value}
    assert ("Prepared by", "Maya Manager") in cells and ("Position", "Back of House Manager") in cells
    summary = client.get("/api/reports/daily/export?fmt=csv", headers=login(MGR)).content.decode("utf-8-sig").splitlines()
    assert "Prepared by,Maya Manager" in summary and "Position,Back of House Manager" in summary


def rec(client, h, kind, **params):
    query = "&".join(f"{k}={v}" for k, v in {"date_from": "2026-03-02", "date_to": "2026-03-02", **params}.items())
    return client.get(f"{URL}/{kind}?{query}", headers=h)


def test_the_monitoring_record_filters_by_result_and_approval(client, login, db, clock):
    seed_readings(client, login)
    m = login(MGR)

    def checks(status):
        doc = monitoring(client, m, location_id=loc_id(db, "MK"), status=status).json()
        return [r["check"] for g in doc["groups"] for r in g["rows"]], len(doc["groups"])

    assert checks("failed") == (["Final rinse"], 1)
    assert checks("passed") == (["Wash"], 1)
    assert checks("awaiting") == (["Final rinse"], 1)
    assert monitoring(client, m, status="nonsense").status_code == 400


def test_the_attendance_record_filters_by_status(client, login, db, clock):
    client.post("/api/attendance/check-in", json={"location_id": loc_id(db, "MK")}, headers=login(STEWARD))
    mgr = login(MGR)
    emp = next(e["id"] for e in client.get("/api/employees?size=50", headers=mgr).json()["items"] if e["employee_code"] == "E005")
    client.post("/api/absences", json={"employee_id": emp, "location_id": loc_id(db, "BQ"), "day": "2026-03-02", "absence_type": "sick_leave"}, headers=mgr)
    clock.advance(hours=10)

    def people(status):
        doc = rec(client, mgr, "attendance", status=status).json()
        return {r["employee"] for g in doc["groups"] for r in g["rows"]}

    assert people("leave") == {"Bina Steward"}
    assert "Sid Steward" in people("present") and "Bina Steward" not in people("present")
    assert "Sid Steward" not in people("absent") and "Bina Steward" not in people("absent")
    assert rec(client, mgr, "attendance", status="nonsense").status_code == 400


def test_the_lot_register_lists_received_lots_and_filters_by_group_supplier_and_status(client, login, db, clock):
    m, sup = login(MGR), login(SUP)
    chicken, salmon = food(client, m), food(client, m, "Salmon fillet", "seafood")
    other = client.post("/api/suppliers", headers=m, json={"name": "Gulf Seafood"}).json()["id"]
    receive(client, sup, db, chicken, supplier_lot="B1", qty=20, unit="kg")
    receive(client, sup, db, salmon, supplier_lot="S1", temp_c=9, supplier_id=other)

    def codes(**params):
        doc = rec(client, m, "lots", **params).json()
        return [r["lot"] for g in doc["groups"] for r in g["rows"]]

    doc = rec(client, m, "lots").json()
    assert doc["title"] == "Lot Register" and codes() == ["B1", "S1"]
    first, second = doc["groups"][0]["rows"]
    assert (first["food"], first["quantity"], first["status"]) == ("Chicken breast", "20.00 kg", "Active")
    assert (second["supplier"], second["arrival_temp"], second["status"]) == ("Gulf Seafood", "9.0 C", "On hold")
    assert dict(doc["groups"][0]["totals"]) == {"Lots": 2, "Active": 1, "Used up": 0, "Discarded": 0, "On hold": 1}
    assert codes(food_group="seafood") == ["S1"] and codes(status="on_hold") == ["S1"] and codes(supplier_id=other) == ["S1"]
    assert codes(food_group="seafood", status="active") == []
    assert rec(client, m, "lots", food_group="nonsense").status_code == 400
    assert rec(client, m, "lots", status="nonsense").status_code == 400
    assert rec(client, m, "receipts", status="active").status_code == 400
    assert b"B1" in client.get(f"{URL}/lots/export?fmt=csv&date_from=2026-03-02&date_to=2026-03-02", headers=m).content


def test_the_monitoring_record_filters_by_point_type_and_one_point_across_locations(client, login, db, clock):
    m = login(MGR)
    make(client, m, db, "cooling")
    make(client, m, db, "cooling", "BQ")
    make(client, m, db, "defrosting")

    def groups(**params):
        return monitoring(client, m, **params).json()["groups"]

    cooling = groups(equipment_kind="Cooling")
    assert len(cooling) == 2 and {dict(g["meta"])["Type"] for g in cooling} == {"Cooling"}
    assert {dict(g["meta"])["Location"] for g in cooling} == {"Main Kitchen", "Banquet Hall"}
    assert {dict(g["meta"])["Type"] for g in groups(equipment_kind="Cooling", location_id=loc_id(db, "BQ"))} == {"Cooling"}
    assert len(groups(equipment_kind="Cooling", location_id=loc_id(db, "BQ"))) == 1
    assert groups(equipment_kind="Nothing like this") == []
    point = next(e for e in client.get("/api/equipment", headers=m).json() if e["kind"] == "Defrosting")
    only = groups(equipment_id=point["id"])
    assert len(only) == 1 and dict(only[0]["meta"])["Monitoring point"].endswith(f"({point['code']})")
    assert len(groups()) > 4


def test_the_checklist_record_filters_by_checklist(client, login, db, clock):
    generate(client, login, db)
    m = login(MGR)
    templates = client.get(f"/api/checklists/templates?location_id={loc_id(db, 'MK')}", headers=m).json()
    daily = next(t for t in templates if t["frequency"] == "daily")
    everything = get(client, m, "checklists", location_id=loc_id(db, "MK")).json()["groups"]
    only = get(client, m, "checklists", location_id=loc_id(db, "MK"), template_id=daily["id"]).json()["groups"]
    assert len(everything) > len(only) >= 1
    assert all(daily["name"] in g["title"] for g in only)
    assert get(client, m, "checklists", template_id=999999).json()["groups"] == []


def test_the_checklist_chart_filters_by_checklist(client, login, db, clock):
    generate(client, login, db)
    m = login(MGR)
    templates = client.get(f"/api/checklists/templates?location_id={loc_id(db, 'MK')}", headers=m).json()
    daily = next(t for t in templates if t["frequency"] == "daily")

    def titles(**params):
        return [g["title"] for g in get(client, m, "checklist_chart", location_id=loc_id(db, "MK"), **params).json()["groups"]]

    everything, only = titles(), titles(template_id=daily["id"])
    assert len(everything) > len(only) >= 1 and all(daily["name"] in t for t in only)
    assert titles(template_id=999999) == []
